import io
import os
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable, List, Optional

from flask import Flask, render_template, request, send_file, jsonify

try:
    import pdfplumber
except ImportError:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None
    get_column_letter = None


ALLOWED_EXTENSIONS = {"pdf", "rpt", "txt"}


@dataclass
class StatementRow:
    date: Optional[str]
    particulars: str
    reference: Optional[str]
    withdrawals: Optional[float]
    deposits: Optional[float]
    balance: Optional[float]
    balance_type: Optional[str]


DATE_PATTERNS = [
    re.compile(r"\b(\d{2}[/-]\d{2}[/-]\d{4})\b"),
    re.compile(r"\b(\d{2}-[A-Za-z]{3}-\d{4})\b"),
    re.compile(r"\b(\d{4}-\d{2}-\d{2})\b"),
]
AMOUNT_PATTERN = re.compile(r"(?<!\d)(\d{1,3}(?:,\d{3})*(?:\.\d{2})?)")
BALANCE_SUFFIX_PATTERN = re.compile(r"\b(Dr|Cr)\b", re.IGNORECASE)


app = Flask(__name__)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/parse", methods=["POST"])
def parse_statement():
    if "statement" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["statement"]
    if file.filename == "":
        return jsonify({"error": "No file selected"}), 400

    if not allowed_file(file.filename):
        return jsonify({"error": "Unsupported file type"}), 400

    raw_text = extract_text(file)
    if not raw_text:
        return jsonify({"error": "Unable to extract text from file"}), 400

    rows = parse_lines(raw_text.splitlines())
    if not rows:
        return jsonify({"error": "No statement rows detected"}), 400

    if Workbook is None:
        return jsonify({"error": "openpyxl is required to generate Excel output"}), 500

    output = build_excel(rows)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="statement.xlsx",
    )


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text(file_storage) -> str:
    extension = file_storage.filename.rsplit(".", 1)[1].lower()
    if extension == "pdf":
        if pdfplumber is None:
            return ""
        return extract_pdf_text(file_storage)
    return file_storage.stream.read().decode("utf-8", errors="ignore")


def extract_pdf_text(file_storage) -> str:
    buffer = io.BytesIO(file_storage.read())
    text_chunks: List[str] = []
    with pdfplumber.open(buffer) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text:
                text_chunks.append(text)
    return "\n".join(text_chunks)


def parse_lines(lines: Iterable[str]) -> List[StatementRow]:
    rows: List[StatementRow] = []
    current_date: Optional[str] = None
    for raw_line in lines:
        line = sanitize_line(raw_line)
        if not line:
            continue

        detected_date = detect_date(line)
        if detected_date:
            current_date = detected_date

        amounts = extract_amounts(line)
        balance_type = detect_balance_type(line)

        row = classify_row(current_date, line, amounts, balance_type)
        if row:
            rows.append(row)
    return merge_multiline_particulars(rows)


def sanitize_line(line: str) -> str:
    return " ".join(line.replace("\t", " ").split())


def detect_date(line: str) -> Optional[str]:
    for pattern in DATE_PATTERNS:
        match = pattern.search(line)
        if match:
            return normalize_date(match.group(1))
    return None


def normalize_date(value: str) -> str:
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return value


def extract_amounts(line: str) -> List[float]:
    amounts = []
    for match in AMOUNT_PATTERN.findall(line):
        try:
            amounts.append(float(match.replace(",", "")))
        except ValueError:
            continue
    return amounts


def detect_balance_type(line: str) -> Optional[str]:
    match = BALANCE_SUFFIX_PATTERN.search(line)
    if match:
        return match.group(1).upper()
    return None


def classify_row(
    current_date: Optional[str],
    line: str,
    amounts: List[float],
    balance_type: Optional[str],
) -> Optional[StatementRow]:
    if current_date is None and not amounts:
        return None

    withdrawals = None
    deposits = None
    balance = None
    reference = None

    if len(amounts) >= 3:
        withdrawals = amounts[-3]
        deposits = amounts[-2]
        balance = amounts[-1]
    elif len(amounts) == 2:
        deposits = amounts[-2]
        balance = amounts[-1]
    elif len(amounts) == 1:
        balance = amounts[0]

    reference_match = re.search(r"(CHQ|REF|UPI|IMPS|NEFT|RTGS)[^\s]*", line, re.IGNORECASE)
    if reference_match:
        reference = reference_match.group(0).upper()

    particulars = strip_amounts(line)
    return StatementRow(
        date=current_date,
        particulars=particulars,
        reference=reference,
        withdrawals=withdrawals,
        deposits=deposits,
        balance=balance,
        balance_type=balance_type,
    )


def strip_amounts(line: str) -> str:
    cleaned = AMOUNT_PATTERN.sub("", line)
    cleaned = BALANCE_SUFFIX_PATTERN.sub("", cleaned)
    return " ".join(cleaned.split())


def merge_multiline_particulars(rows: List[StatementRow]) -> List[StatementRow]:
    merged: List[StatementRow] = []
    for row in rows:
        if merged and row.date == merged[-1].date and row.particulars and not row.withdrawals and not row.deposits:
            merged[-1].particulars = f"{merged[-1].particulars} {row.particulars}".strip()
        else:
            merged.append(row)
    return merged


def build_excel(rows: List[StatementRow]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Statement"

    headers = [
        "Date",
        "Particulars",
        "Reference",
        "Withdrawals",
        "Deposits",
        "Balance",
        "Balance Type",
    ]
    sheet.append(headers)

    for row in rows:
        sheet.append(
            [
                row.date,
                row.particulars,
                row.reference,
                row.withdrawals,
                row.deposits,
                row.balance,
                row.balance_type,
            ]
        )

    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        column_letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[column_letter].width = min(length + 2, 40)

    output = io.BytesIO()
    workbook.save(output)
    return output


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
